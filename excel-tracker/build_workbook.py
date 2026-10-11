#!/usr/bin/env python3
"""Builds the APD Project Tracker Excel workbook.

Usage:  python3 build_workbook.py [output_dir]

Produces two files: APD_Project_Tracker.xlsx (one example record per register,
to show the expected format) and APD_Project_Tracker_BLANK.xlsx (no data).
Everything the workbook shows is a live formula over the register sheets, so
adding or editing rows recalculates the whole dashboard.
"""
import datetime as dt
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

FONT = "Arial"
NAVY = "1F3864"
GREY = "F2F2F2"
CALC_FILL = PatternFill("solid", fgColor="EDEDED")
HDR_FILL = PatternFill("solid", fgColor=NAVY)
HDR_CALC_FILL = PatternFill("solid", fgColor="7F7F7F")
INPUT_FONT = Font(name=FONT, size=10, color="0000FF")
CALC_FONT = Font(name=FONT, size=10, color="000000")
HDR_FONT = Font(name=FONT, size=10, bold=True, color="FFFFFF")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
GREEN_FILL = PatternFill("solid", bgColor="C6EFCE", fgColor="C6EFCE")
AMBER_FILL = PatternFill("solid", bgColor="FFEB9C", fgColor="FFEB9C")
RED_FILL = PatternFill("solid", bgColor="FFC7CE", fgColor="FFC7CE")
GREY_FILL = PatternFill("solid", bgColor="D9D9D9", fgColor="D9D9D9")

DATE_FMT = "dd-mmm-yyyy"
MONEY_FMT = '$#,##0;($#,##0);"-"'
PCT_FMT = "0%"

EX_PROJ = "Pump Station Refurbishment (example)"
TODAY = dt.date.today()
SMALL = False


def d(offset_days):
    return TODAY + dt.timedelta(days=offset_days)


# ---------------------------------------------------------------------------
# Dropdown lists (display labels, as shown in the web app)
# ---------------------------------------------------------------------------
LISTS = {
    "ProjectStatus": ["Planning", "Design", "Tender", "Construction", "Project Completion",
                      "Defects Liability", "Complete", "On Hold", "Cancelled"],
    "TaskStatus": ["To Do", "In Progress", "Done"],
    "Recurrence": ["None", "Daily", "Weekly", "Monthly"],
    "CashType": ["Actual", "Forecast"],
    "VariationStatus": ["Proposed", "Approved", "Rejected"],
    "RiskType": ["Risk", "Opportunity"],
    "Likelihood": ["Rare", "Unlikely", "Possible", "Likely", "Almost Certain"],
    "Consequence": ["Insignificant", "Minor", "Moderate", "Major", "Severe"],
    "RiskStatus": ["Open", "Mitigating", "Closed"],
    "ActionType": ["Action", "Decision"],
    "ActionStatus": ["Open", "In Progress", "Done"],
    "PermitStatus": ["Pending", "Approved", "Rejected"],
    "LessonType": ["Positive", "Negative"],
    "SqeCategory": ["Safety", "Quality", "Environment"],
    "SqeSeverity": ["Low", "Medium", "High", "Critical"],
    "SqeStatus": ["Open", "In Progress", "Closed"],
    "EotStatus": ["Submitted", "Approved", "Rejected"],
    "ChecklistPhase": ["Commissioning", "Handover"],
    "ChecklistStatus": ["Not Started", "In Progress", "Complete", "Not Applicable"],
    "DefectStatus": ["Open", "Rectified", "Verified", "Disputed"],
    "PaymentStatus": ["Submitted", "Certified", "Disputed", "Paid"],
    "IWType": ["Insurance", "Warranty"],
    "RfiStatus": ["Open", "Answered", "Closed"],
    "DocStatus": ["Draft", "For Review", "Approved", "Superseded"],
    "EnquiryCategory": ["Noise", "Traffic", "Dust", "Access", "Other"],
    "EnquiryStatus": ["Open", "In Progress", "Resolved", "Closed"],
    "SubStatus": ["Engaged", "Active", "Completed", "Terminated"],
    "PartyType": ["Consultant", "Contractor"],
    "ContractStatus": ["Active", "Complete", "Terminated"],
    "Direction": ["Incoming", "Outgoing"],
    "CorrStatus": ["Open", "Closed"],
    "LimitType": ["Max", "Min"],
    "SafetyInspStatus": ["Open", "In Progress", "Closed"],
    "TenderMethod": ["Open", "Selective", "Direct Engage", "Expression of Interest"],
    "TenderStatus": ["Planned", "Issued", "Evaluation", "Awarded", "Cancelled"],
    "Rating": [1, 2, 3, 4, 5],
}


class Col:
    def __init__(self, h, kind="text", w=16, lst=None, f=None, fmt=None, note=None):
        self.h, self.kind, self.w, self.lst, self.f, self.fmt, self.note = h, kind, w, lst, f, fmt, note
        self.letter = None

    @property
    def is_calc(self):
        return self.f is not None


class Reg:
    def __init__(self, name, key, cols, ex=None, cap=500, blurb=""):
        self.name, self.key, self.cols, self.ex, self.cap, self.blurb = name, key, cols, ex or [], cap, blurb
        for i, c in enumerate(cols, start=1):
            c.letter = L(i)
        self.map = {c.h: c.letter for c in cols}
        self.last = cols[-1].letter


REGS = {}


def reg(name, key, cols, ex=None, cap=500, blurb=""):
    REGS[name] = Reg(name, key, cols, ex, cap, blurb)
    return REGS[name]


def q(name):
    return f"'{name}'"


def R(sheet, h):
    r = REGS[sheet]
    return f"{q(sheet)}!${r.map[h]}$2:${r.map[h]}${r.cap + 1}"


def cref(sheet, h, row):
    return f"{REGS[sheet].map[h]}{row}"


def make_c(sheet, row):
    return lambda h: cref(sheet, h, row)


# ---------------------------------------------------------------------------
# Shared formula fragments
# ---------------------------------------------------------------------------
def P(w=26):
    return Col("Project", "proj", w)


def status_bucket(c, status_h, done_values, due_h):
    done = ",".join(f'{c(status_h)}="{v}"' for v in done_values)
    return (f'IF(OR({done}),"Done",IF({c(due_h)}="","No Due Date",IF({c(due_h)}<TODAY(),"Overdue",'
            f'IF({c(due_h)}=TODAY(),"Due Today",IF({c(due_h)}<TODAY()+7,"Due This Week","Later")))))')


def expiry(c, h):
    return f'IF({c(h)}="","",IF({c(h)}<TODAY(),"Expired",IF({c(h)}<=TODAY()+60,"Expiring Soon","OK")))'


# Projects --------------------------------------------------------------------
EXEMPT = lambda c: (f'OR({c("Status")}="Project Completion",{c("Status")}="Defects Liability",'
                    f'{c("Status")}="Complete",{c("Status")}="Cancelled")')


def pf(body):
    """Wrap a Projects-row formula so blank rows stay blank."""
    return lambda c, r: f'=IF({c("Project Name")}="","",{body(c, r)})'


projects_cols = [
    Col("Project Name", "text", 34, note="Unique name. Registers pick the project from this list."),
    Col("Description", "long", 34),
    Col("Asset Type", "text", 18),
    Col("Status", "list", 18, "ProjectStatus"),
    Col("BP Ref No", "text", 12),
    Col("Project Manager", "text", 18),
    Col("Project Director", "text", 18),
    Col("Project Sponsor", "text", 18),
    Col("Location", "text", 18),
    Col("Contractor", "text", 20),
    Col("Contract Value", "money", 15, note="Display only (not part of health maths)."),
    Col("Original Budget", "money", 15),
    Col("Latest Approved Budget", "money", 18, note="Leave blank if unchanged from the original."),
    Col("Start Date", "date", 14),
    Col("Target Completion", "date", 14),
    Col("Actual Completion", "date", 14),
    # computed ---------------------------------------------------------------
    Col("Approved Variations", "calc", 15, f=pf(lambda c, r: (
        f'SUMIFS({R("Variations", "Cost Impact")},{R("Variations", "Project")},{c("Project Name")},'
        f'{R("Variations", "Status")},"Approved")')), fmt=MONEY_FMT),
    Col("Effective Budget", "calc", 15, f=pf(lambda c, r: (
        f'IF({c("Latest Approved Budget")}<>"",{c("Latest Approved Budget")},'
        f'IF({c("Original Budget")}="","",{c("Original Budget")}+{c("Approved Variations")}))')), fmt=MONEY_FMT),
    Col("Expenditure to Date", "calc", 15, f=pf(lambda c, r: (
        f'SUMIFS({R("Cash Flow", "Amount")},{R("Cash Flow", "Project")},{c("Project Name")},'
        f'{R("Cash Flow", "Type")},"Actual")')), fmt=MONEY_FMT),
    Col("Forecast at Completion", "calc", 16, f=pf(lambda c, r: (
        f'SUMIFS({R("Cash Flow", "Counted Amount")},{R("Cash Flow", "Project")},{c("Project Name")})')),
        fmt=MONEY_FMT, note="Actual for months with an actual, otherwise forecast (no double counting)."),
    Col("Budget Variance", "calc", 15, f=pf(lambda c, r: (
        f'IF({c("Effective Budget")}="","",{c("Effective Budget")}-{c("Forecast at Completion")})')), fmt=MONEY_FMT),
    Col("Approved EOT Days", "calc", 12, f=pf(lambda c, r: (
        f'SUMIFS({R("EOT Claims", "Days Counted")},{R("EOT Claims", "Project")},{c("Project Name")})'))),
    Col("Revised Completion", "calc", 15, f=pf(lambda c, r: (
        f'IF({c("Target Completion")}="","",{c("Target Completion")}+{c("Approved EOT Days")})')), fmt=DATE_FMT),
    Col("Overdue", "calc", 11, f=pf(lambda c, r: (
        f'IF({c("Target Completion")}="","",IF(AND(NOT({EXEMPT(c)}),{c("Target Completion")}<TODAY()),"Overdue",""))'))),
    Col("Schedule Elapsed %", "calc", 13, f=pf(lambda c, r: (
        f'IF(OR({c("Start Date")}="",{c("Revised Completion")}=""),"",'
        f'IF({c("Revised Completion")}<={c("Start Date")},"",(TODAY()-{c("Start Date")})/'
        f'({c("Revised Completion")}-{c("Start Date")})))')), fmt=PCT_FMT),
    Col("Schedule Health", "calc", 12, f=pf(lambda c, r: (
        f'IF(OR({c("Schedule Elapsed %")}="",{EXEMPT(c)}),"Green",IF({c("Schedule Elapsed %")}>1.15,"Red",'
        f'IF({c("Schedule Elapsed %")}>=0.95,"Amber","Green")))'))),
    Col("Forecast % of Budget", "calc", 13, f=pf(lambda c, r: (
        f'IF({c("Effective Budget")}="","",IF(OR({c("Effective Budget")}<=0,'
        f'COUNTIFS({R("Cash Flow", "Project")},{c("Project Name")})=0),"",'
        f'{c("Forecast at Completion")}/{c("Effective Budget")}))')), fmt=PCT_FMT),
    Col("Budget Health", "calc", 12, f=pf(lambda c, r: (
        f'IF({c("Forecast % of Budget")}="","Green",IF({c("Forecast % of Budget")}>1.15,"Red",'
        f'IF({c("Forecast % of Budget")}>=0.95,"Amber","Green")))'))),
    Col("Critical Risks", "calc", 10, f=pf(lambda c, r: (
        f'COUNTIFS({R("Risks", "Project")},{c("Project Name")},{R("Risks", "Type")},"Risk",'
        f'{R("Risks", "Status")},"<>Closed",{R("Risks", "Inherent Band")},"Critical")'))),
    Col("High Risks", "calc", 10, f=pf(lambda c, r: (
        f'COUNTIFS({R("Risks", "Project")},{c("Project Name")},{R("Risks", "Type")},"Risk",'
        f'{R("Risks", "Status")},"<>Closed",{R("Risks", "Inherent Band")},"High")'))),
    Col("Risk Health", "calc", 11, f=pf(lambda c, r: (
        f'IF({c("Critical Risks")}>0,"Red",IF({c("High Risks")}>0,"Amber","Green"))'))),
    Col("Overall Health", "calc", 12, f=pf(lambda c, r: (
        f'IF(OR({c("Schedule Health")}="Red",{c("Budget Health")}="Red",{c("Risk Health")}="Red"),"Red",'
        f'IF(OR({c("Schedule Health")}="Amber",{c("Budget Health")}="Amber",{c("Risk Health")}="Amber"),"Amber","Green"))'))),
    Col("Open Tasks", "calc", 9, f=pf(lambda c, r: (
        f'COUNTIFS({R("Tasks", "Project")},{c("Project Name")},{R("Tasks", "Status")},"<>Done",{R("Tasks", "Title")},"<>")'))),
    Col("Open Actions", "calc", 9, f=pf(lambda c, r: (
        f'COUNTIFS({R("Actions", "Project")},{c("Project Name")},{R("Actions", "Status")},"<>Done",{R("Actions", "Title")},"<>")'))),
    Col("Open RFIs", "calc", 9, f=pf(lambda c, r: (
        f'COUNTIFS({R("RFIs", "Project")},{c("Project Name")},{R("RFIs", "Status")},"Open")'))),
    Col("Open Defects", "calc", 9, f=pf(lambda c, r: (
        f'COUNTIFS({R("Defects", "Project")},{c("Project Name")},{R("Defects", "Status")},"Open")'))),
    Col("Health Notes", "calc", 60, f=pf(lambda c, r: (
        f'"Schedule: "&IF({c("Schedule Elapsed %")}="","no dates / not tracked",TEXT({c("Schedule Elapsed %")},"0%")&" of timeframe elapsed")'
        f'&" | Budget: "&IF({c("Forecast % of Budget")}="","no approved budget or cash flow",TEXT({c("Forecast % of Budget")},"0%")&" of budget forecast")'
        f'&" | Risk: "&{c("Critical Risks")}&" critical, "&{c("High Risks")}&" high open"'))),
]
reg("Projects", "Project Name", projects_cols, cap=100, blurb="One row per project. Grey columns are formulas - do not type in them.", ex=[{
    "Project Name": EX_PROJ, "Description": "Replace two pumps and switchboard at the Nathalia pump station",
    "Asset Type": "Pump Station", "Status": "Construction", "BP Ref No": "3323", "Project Manager": "J. Smith",
    "Project Director": "A. Jones", "Project Sponsor": "Operations", "Location": "Nathalia",
    "Contractor": "ABC Civil Pty Ltd", "Contract Value": 920000, "Original Budget": 1000000,
    "Latest Approved Budget": 1060000, "Start Date": d(-150), "Target Completion": d(60)}])

# Tasks -----------------------------------------------------------------------
reg("Tasks", "Title", [
    Col("Project", "proj", 26, note="Leave blank for a portfolio-wide task."),
    Col("Title", "text", 36), Col("Description", "long", 34),
    Col("Status", "list", 12, "TaskStatus"), Col("Due Date", "date", 13),
    Col("Recurrence", "list", 12, "Recurrence", note="Marking a recurring task Done creates the next one (macro)."),
    Col("Completed Date", "date", 13),
    Col("Bucket", "calc", 14, f=lambda c, r: f'=IF({c("Title")}="","",{status_bucket(c, "Status", ["Done"], "Due Date")})'),
    Col("Upcoming Rank", "calc", 12, f=lambda c, r: (
        f'=IF(AND({c("Title")}<>"",{c("Status")}<>"Done",{c("Due Date")}<>""),{c("Due Date")}+ROW()/100000,"")'), fmt="0.00000"),
], ex=[{"Project": EX_PROJ, "Title": "Order replacement pump seals", "Description": "Confirm lead time with supplier",
        "Status": "In Progress", "Due Date": d(5), "Recurrence": "None"}])

# Actions ---------------------------------------------------------------------
reg("Actions", "Title", [
    P(), Col("Type", "list", 11, "ActionType"), Col("Title", "text", 36), Col("Description", "long", 34),
    Col("Category", "text", 14), Col("Owner", "text", 16), Col("Due Date", "date", 13),
    Col("Status", "list", 12, "ActionStatus"), Col("Date Raised", "date", 13), Col("Date Closed", "date", 13),
    Col("Meeting", "text", 28, note="Optional: title of the meeting the action came from."),
    Col("Notes", "long", 30),
    Col("Bucket", "calc", 14, f=lambda c, r: f'=IF({c("Title")}="","",{status_bucket(c, "Status", ["Done"], "Due Date")})'),
    Col("Upcoming Rank", "calc", 12, f=lambda c, r: (
        f'=IF(AND({c("Title")}<>"",{c("Status")}<>"Done",{c("Due Date")}<>""),{c("Due Date")}+ROW()/100000,"")'), fmt="0.00000"),
    Col("Summary Seq", "calc", 10, f=lambda c, r: (
        f'=IF(AND({c("Project")}<>"",{c("Project")}=\'Project Summary\'!$C$3,{c("Status")}<>"Done",{c("Title")}<>""),'
        f'COUNTIFS($A$2:$A{r},\'Project Summary\'!$C$3,$H$2:$H{r},"<>Done",$C$2:$C{r},"<>"),"")')),
], ex=[{"Project": EX_PROJ, "Type": "Action", "Title": "Issue revised switchboard drawings", "Category": "Design",
        "Owner": "M. Lee", "Due Date": d(-2), "Status": "Open", "Date Raised": d(-20)}])

# Risks -----------------------------------------------------------------------
def score(c, lik, con):
    return f'MATCH({c(lik)},lst_Likelihood,0)*MATCH({c(con)},lst_Consequence,0)'


def band(s):
    return f'IF({s}>=16,"Critical",IF({s}>=10,"High",IF({s}>=5,"Medium","Low")))'


reg("Risks", "Title", [
    P(), Col("Type", "list", 12, "RiskType"), Col("Title", "text", 34), Col("Description", "long", 32),
    Col("Category", "text", 14),
    Col("Inherent Likelihood", "list", 15, "Likelihood"), Col("Inherent Consequence", "list", 15, "Consequence"),
    Col("Residual Likelihood", "list", 15, "Likelihood"), Col("Residual Consequence", "list", 15, "Consequence"),
    Col("Mitigation", "long", 30), Col("Owner", "text", 16), Col("Due Date", "date", 13),
    Col("Status", "list", 12, "RiskStatus"),
    Col("Inherent Score", "calc", 10, f=lambda c, r: (
        f'=IF(OR({c("Title")}="",{c("Inherent Likelihood")}="",{c("Inherent Consequence")}=""),"",{score(c, "Inherent Likelihood", "Inherent Consequence")})')),
    Col("Inherent Band", "calc", 12, f=lambda c, r: f'=IF({c("Inherent Score")}="","",{band(c("Inherent Score"))})'),
    Col("Residual Score", "calc", 10, f=lambda c, r: (
        f'=IF(OR({c("Title")}="",{c("Residual Likelihood")}="",{c("Residual Consequence")}=""),"",{score(c, "Residual Likelihood", "Residual Consequence")})')),
    Col("Residual Band", "calc", 12, f=lambda c, r: f'=IF({c("Residual Score")}="","",{band(c("Residual Score"))})'),
    Col("Overdue", "calc", 10, f=lambda c, r: (
        f'=IF(AND({c("Title")}<>"",{c("Status")}<>"Closed",{c("Due Date")}<>"",{c("Due Date")}<TODAY()),"Overdue","")')),
    Col("Top Rank", "calc", 10, f=lambda c, r: (
        f'=IF(AND({c("Inherent Score")}<>"",{c("Type")}="Risk",{c("Status")}<>"Closed"),{c("Inherent Score")}+(1-ROW()/100000),"")'), fmt="0.00000"),
    Col("Summary Rank", "calc", 10, f=lambda c, r: (
        f'=IF(AND({c("Top Rank")}<>"",{c("Project")}=\'Project Summary\'!$C$3),{c("Top Rank")},"")'), fmt="0.00000"),
], ex=[{"Project": EX_PROJ, "Type": "Risk", "Title": "Switchboard delivery delay", "Category": "Programme",
        "Inherent Likelihood": "Likely", "Inherent Consequence": "Major", "Residual Likelihood": "Possible",
        "Residual Consequence": "Moderate", "Mitigation": "Expedite order; confirm factory slot", "Owner": "J. Smith",
        "Due Date": d(14), "Status": "Mitigating"}])

# RFIs ------------------------------------------------------------------------
reg("RFIs", "RFI Number", [
    P(), Col("RFI Number", "text", 13), Col("Subject", "text", 30), Col("Question", "long", 36),
    Col("Response", "long", 30), Col("Raised To", "text", 16), Col("Status", "list", 11, "RfiStatus"),
    Col("Date Raised", "date", 13), Col("Date Required", "date", 13), Col("Date Answered", "date", 13),
    Col("Notes", "long", 28),
    Col("Flag", "calc", 12, f=lambda c, r: (
        f'=IF(OR({c("RFI Number")}="",{c("Status")}<>"Open",{c("Date Required")}=""),"",'
        f'IF({c("Date Required")}<TODAY(),"Overdue",IF({c("Date Required")}<=TODAY()+7,"Due Soon","")))')),
], ex=[{"Project": EX_PROJ, "RFI Number": "RFI-001", "Subject": "Pump base level", "Question": "Confirm finished level of plinth",
        "Raised To": "Design consultant", "Status": "Open", "Date Raised": d(-6), "Date Required": d(3)}])

# Defects ---------------------------------------------------------------------
reg("Defects", "Title", [
    P(), Col("Title", "text", 32), Col("Description", "long", 32), Col("Location", "text", 18), Col("Category", "text", 14),
    Col("Assigned To", "text", 16), Col("Due Date", "date", 13), Col("Status", "list", 11, "DefectStatus"),
    Col("Date Raised", "date", 13), Col("Date Closed", "date", 13), Col("Retention Held", "money", 14),
    Col("Notes", "long", 28),
    Col("Flag", "calc", 12, f=lambda c, r: (
        f'=IF(OR({c("Title")}="",{c("Status")}<>"Open",{c("Due Date")}=""),"",IF({c("Due Date")}<TODAY(),"Overdue",""))')),
], ex=[{"Project": EX_PROJ, "Title": "Chipped paint on pump housing", "Location": "Pump 2", "Category": "Finishes",
        "Assigned To": "ABC Civil", "Due Date": d(10), "Status": "Open", "Date Raised": d(-3), "Retention Held": 2500}])

# Variations ------------------------------------------------------------------
reg("Variations", "Title", [
    P(), Col("Reference", "text", 12), Col("Title", "text", 32), Col("Description", "long", 32), Col("Category", "text", 14),
    Col("Cost Impact", "money", 14), Col("Time Impact Days", "int", 12), Col("Status", "list", 12, "VariationStatus"),
    Col("Date Raised", "date", 13), Col("Date Decided", "date", 13), Col("Notes", "long", 28),
    Col("Summary Seq", "calc", 10, f=lambda c, r: (
        f'=IF(AND({c("Project")}<>"",{c("Project")}=\'Project Summary\'!$C$3,{c("Title")}<>""),'
        f'COUNTIFS($A$2:$A{r},\'Project Summary\'!$C$3,$C$2:$C{r},"<>"),"")')),
], ex=[{"Project": EX_PROJ, "Reference": "VO-01", "Title": "Additional valve chamber", "Category": "Scope",
        "Cost Impact": 60000, "Time Impact Days": 10, "Status": "Approved", "Date Raised": d(-40), "Date Decided": d(-30)}])

# EOT -------------------------------------------------------------------------
reg("EOT Claims", "Title", [
    P(), Col("Reference", "text", 12), Col("Title", "text", 32), Col("Cause", "text", 18), Col("Description", "long", 32),
    Col("Days Claimed", "int", 12), Col("Days Granted", "int", 12), Col("Status", "list", 12, "EotStatus"),
    Col("Date Submitted", "date", 13), Col("Date Decided", "date", 13), Col("Notes", "long", 28),
    Col("Days Counted", "calc", 12, f=lambda c, r: (
        f'=IF(OR({c("Title")}="",{c("Status")}<>"Approved"),0,IF({c("Days Granted")}<>"",{c("Days Granted")},N({c("Days Claimed")})))')),
], ex=[{"Project": EX_PROJ, "Reference": "EOT-01", "Title": "Wet weather delay", "Cause": "Weather",
        "Days Claimed": 8, "Days Granted": 5, "Status": "Approved", "Date Submitted": d(-25), "Date Decided": d(-15)}])

# Cash flow -------------------------------------------------------------------
reg("Cash Flow", "Month", [
    P(), Col("Month", "date", 13, note="Any date in the month; it is normalised to the 1st."),
    Col("Type", "list", 11, "CashType"), Col("Amount", "money", 14),
    Col("Month Key", "calc", 13, f=lambda c, r: f'=IF(OR({c("Project")}="",{c("Month")}=""),"",DATE(YEAR({c("Month")}),MONTH({c("Month")}),1))', fmt="mmm-yyyy"),
    Col("Counted Amount", "calc", 14, f=lambda c, r: (
        f'=IF(OR({c("Month Key")}="",{c("Amount")}=""),0,IF({c("Type")}="Actual",{c("Amount")},'
        f'IF(COUNTIFS($A$2:$A$2001,{c("Project")},$E$2:$E$2001,{c("Month Key")},$C$2:$C$2001,"Actual")>0,0,{c("Amount")})))'), fmt=MONEY_FMT,
        note="An Actual replaces that month's Forecast so nothing is double counted."),
], cap=2000, ex=[
    {"Project": EX_PROJ, "Month": d(-60), "Type": "Actual", "Amount": 250000},
    {"Project": EX_PROJ, "Month": d(-30), "Type": "Actual", "Amount": 310000},
    {"Project": EX_PROJ, "Month": d(0), "Type": "Forecast", "Amount": 280000},
    {"Project": EX_PROJ, "Month": d(30), "Type": "Forecast", "Amount": 220000}])

# Milestones ------------------------------------------------------------------
reg("Milestones", "Task / Activity", [
    P(), Col("WBS Item", "text", 12), Col("Task / Activity", "text", 34), Col("% Complete", "int", 11),
    Col("Start Date", "date", 13), Col("Original Completion", "date", 14), Col("Forecast Completion", "date", 14),
    Col("Actual Completion", "date", 14), Col("Notes", "long", 28),
    Col("Slip (days)", "calc", 10, f=lambda c, r: (
        f'=IF(OR({c("Task / Activity")}="",{c("Original Completion")}="",{c("Forecast Completion")}=""),"",{c("Forecast Completion")}-{c("Original Completion")})')),
    Col("Summary Seq", "calc", 10, f=lambda c, r: (
        f'=IF(AND({c("Project")}<>"",{c("Project")}=\'Project Summary\'!$C$3,{c("Task / Activity")}<>""),'
        f'COUNTIFS($A$2:$A{r},\'Project Summary\'!$C$3,$C$2:$C{r},"<>"),"")')),
], ex=[
    {"Project": EX_PROJ, "WBS Item": "1.1", "Task / Activity": "Procure pumps and switchboard", "% Complete": 80,
     "Start Date": d(-120), "Original Completion": d(-10), "Forecast Completion": d(10)},
    {"Project": EX_PROJ, "WBS Item": "1.2", "Task / Activity": "Install and commission", "% Complete": 20,
     "Start Date": d(-30), "Original Completion": d(50), "Forecast Completion": d(60)}])

# Payment claims --------------------------------------------------------------
reg("Payment Claims", "Claim Number", [
    P(), Col("Contract", "text", 24, note="Party name of the contract this claim is against (optional)."),
    Col("Claim Number", "text", 13), Col("Period Ending", "date", 13), Col("Claimed Amount", "money", 14),
    Col("Certified Amount", "money", 14), Col("Status", "list", 12, "PaymentStatus"), Col("Date Submitted", "date", 13),
    Col("Date Certified", "date", 13), Col("Date Paid", "date", 13), Col("Notes", "long", 28),
], ex=[{"Project": EX_PROJ, "Contract": "ABC Civil Pty Ltd", "Claim Number": "PC-01", "Period Ending": d(-31),
        "Claimed Amount": 250000, "Certified Amount": 245000, "Status": "Paid", "Date Submitted": d(-28),
        "Date Certified": d(-20), "Date Paid": d(-8)}])

# Contracts -------------------------------------------------------------------
reg("Contracts", "Party Name", [
    P(), Col("Party Type", "list", 12, "PartyType"), Col("Party Name", "text", 26), Col("Discipline / Scope", "text", 20),
    Col("Contract Number", "text", 14), Col("Contract Form", "text", 14), Col("Contract Value", "money", 14),
    Col("Contingency", "money", 13), Col("PO Value", "money", 13), Col("Provisional Items", "money", 14),
    Col("Retention %", "int", 10), Col("DLP Months", "int", 10), Col("Execution Date", "date", 13),
    Col("Start Date", "date", 13), Col("Completion Date", "date", 13), Col("Status", "list", 12, "ContractStatus"),
    Col("Notes", "long", 28),
    Col("Claimed to Date", "calc", 14, f=lambda c, r: (
        f'=IF({c("Party Name")}="","",SUMIFS({R("Payment Claims", "Claimed Amount")},{R("Payment Claims", "Project")},{c("Project")},'
        f'{R("Payment Claims", "Contract")},{c("Party Name")},{R("Payment Claims", "Status")},"<>Disputed"))'), fmt=MONEY_FMT),
    Col("Remaining", "calc", 14, f=lambda c, r: (
        f'=IF(OR({c("Party Name")}="",{c("Contract Value")}=""),"",{c("Contract Value")}-{c("Claimed to Date")})'), fmt=MONEY_FMT),
], ex=[{"Project": EX_PROJ, "Party Type": "Contractor", "Party Name": "ABC Civil Pty Ltd", "Discipline / Scope": "Construction",
        "Contract Number": "C-1042", "Contract Form": "AS4000-1997", "Contract Value": 920000, "Contingency": 46000,
        "Retention %": 5, "DLP Months": 12, "Execution Date": d(-170), "Start Date": d(-150), "Completion Date": d(60),
        "Status": "Active"}])

# Subcontractors --------------------------------------------------------------
reg("Subcontractors", "Company Name", [
    P(), Col("Company Name", "text", 26), Col("Trade", "text", 16), Col("Contact Name", "text", 18),
    Col("Contact Details", "text", 24), Col("Contract Value", "money", 14), Col("Start Date", "date", 13),
    Col("End Date", "date", 13), Col("Status", "list", 12, "SubStatus"), Col("Insurance Expiry", "date", 14),
    Col("Notes", "long", 28),
    Col("Insurance Flag", "calc", 14, f=lambda c, r: (
        f'=IF(OR({c("Company Name")}="",OR({c("Status")}="Completed",{c("Status")}="Terminated")),"",{expiry(c, "Insurance Expiry")})')),
], ex=[{"Project": EX_PROJ, "Company Name": "Sparky Electrical", "Trade": "Electrical", "Contact Name": "T. Nguyen",
        "Contact Details": "0400 000 000", "Contract Value": 85000, "Start Date": d(-40), "End Date": d(40),
        "Status": "Active", "Insurance Expiry": d(30)}])

# Permits ---------------------------------------------------------------------
reg("Permits", "Title", [
    P(), Col("Title", "text", 30), Col("Category", "text", 14), Col("Reference Number", "text", 16),
    Col("Issuing Authority", "text", 20), Col("Status", "list", 11, "PermitStatus"), Col("Date Applied", "date", 13),
    Col("Date Approved", "date", 13), Col("Expiry Date", "date", 13), Col("Notes", "long", 28),
    Col("Expiry Flag", "calc", 14, f=lambda c, r: f'=IF({c("Title")}="","",{expiry(c, "Expiry Date")})'),
], ex=[{"Project": EX_PROJ, "Title": "Works within waterway permit", "Category": "Environmental", "Reference Number": "WW-4471",
        "Issuing Authority": "Catchment Authority", "Status": "Approved", "Date Applied": d(-90), "Date Approved": d(-60),
        "Expiry Date": d(45)}])

# Insurance & warranties ------------------------------------------------------
reg("Insurance Warranties", "Title", [
    P(), Col("Type", "list", 12, "IWType"), Col("Title", "text", 28), Col("Provider", "text", 20),
    Col("Reference Number", "text", 16), Col("Coverage Details", "long", 30), Col("Start Date", "date", 13),
    Col("Expiry Date", "date", 13), Col("Notes", "long", 28),
    Col("Expiry Flag", "calc", 14, f=lambda c, r: f'=IF({c("Title")}="","",{expiry(c, "Expiry Date")})'),
], ex=[{"Project": EX_PROJ, "Type": "Insurance", "Title": "Public liability", "Provider": "Insurer Co",
        "Reference Number": "PL-99120", "Start Date": d(-170), "Expiry Date": d(190)}])

# Stakeholders ----------------------------------------------------------------
reg("Stakeholders", "Name", [
    P(), Col("Name", "text", 24), Col("Organization", "text", 24), Col("Role", "text", 20), Col("Email", "text", 28),
    Col("Phone", "text", 16), Col("Notes", "long", 30),
], ex=[{"Project": EX_PROJ, "Name": "Pat Citizen", "Organization": "Local council", "Role": "Approvals officer",
        "Email": "pat@example.com", "Phone": "0400 111 222"}])

# Meeting minutes -------------------------------------------------------------
reg("Meeting Minutes", "Title", [
    P(), Col("Title", "text", 30), Col("Meeting Date", "date", 13), Col("Meeting Type", "text", 16), Col("Location", "text", 18),
    Col("Attendees", "long", 28), Col("Summary", "long", 40), Col("Next Meeting", "date", 13),
    Col("Notes", "long", 28), Col("Document Link", "link", 30, note="Paste a OneDrive/SharePoint link to the minutes file."),
], ex=[{"Project": EX_PROJ, "Title": "Fortnightly progress meeting", "Meeting Date": d(-4), "Meeting Type": "Progress",
        "Location": "Site office", "Attendees": "J. Smith, M. Lee, ABC Civil", "Summary": "Pumps on order; switchboard slipped 10 days",
        "Next Meeting": d(10)}])

# Correspondence --------------------------------------------------------------
reg("Correspondence", "Subject", [
    P(), Col("Reference", "text", 12), Col("Direction", "list", 11, "Direction"), Col("Correspondent", "text", 22),
    Col("Subject", "text", 32), Col("Summary", "long", 32), Col("Date Sent", "date", 13), Col("Response Due", "date", 13),
    Col("Status", "list", 10, "CorrStatus"), Col("Date Closed", "date", 13), Col("Notes", "long", 26),
    Col("Flag", "calc", 12, f=lambda c, r: (
        f'=IF(OR({c("Subject")}="",{c("Status")}<>"Open",{c("Response Due")}=""),"",IF({c("Response Due")}<TODAY(),"Overdue",""))')),
], ex=[{"Project": EX_PROJ, "Reference": "COR-01", "Direction": "Incoming", "Correspondent": "ABC Civil",
        "Subject": "Notice of delay - switchboard", "Date Sent": d(-9), "Response Due": d(-1), "Status": "Open"}])

# Documents -------------------------------------------------------------------
reg("Documents", "Title", [
    P(), Col("Document Number", "text", 16), Col("Title", "text", 32), Col("Category", "text", 16), Col("Revision", "text", 9),
    Col("Status", "list", 12, "DocStatus"), Col("Issue Date", "date", 13), Col("Description", "long", 28),
    Col("Notes", "long", 24), Col("File Link", "link", 30, note="OneDrive/SharePoint link to the document."),
], ex=[{"Project": EX_PROJ, "Document Number": "PS-DWG-001", "Title": "General arrangement", "Category": "Drawing",
        "Revision": "B", "Status": "Approved", "Issue Date": d(-45)}])

# SQE -------------------------------------------------------------------------
reg("SQE Issues", "Title", [
    P(), Col("Category", "list", 13, "SqeCategory"), Col("Severity", "list", 11, "SqeSeverity"), Col("Title", "text", 32),
    Col("Description", "long", 32), Col("Corrective Action", "long", 30), Col("Date Occurred", "date", 13),
    Col("Status", "list", 11, "SqeStatus"), Col("Date Closed", "date", 13), Col("Reported By", "text", 16),
], ex=[{"Project": EX_PROJ, "Category": "Safety", "Severity": "Medium", "Title": "Near miss - excavator swing radius",
        "Corrective Action": "Exclusion zone barriers installed", "Date Occurred": d(-12), "Status": "In Progress", "Reported By": "M. Lee"}])

reg("Safety Inspections", "Inspection Type", [
    P(), Col("Inspection Date", "date", 13), Col("Inspection Type", "text", 22), Col("Location", "text", 18),
    Col("Conducted By", "text", 16), Col("Score", "int", 8), Col("Findings", "long", 32), Col("Corrective Action", "long", 30),
    Col("Status", "list", 11, "SafetyInspStatus"), Col("Date Closed", "date", 13), Col("Notes", "long", 24),
], ex=[{"Project": EX_PROJ, "Inspection Date": d(-7), "Inspection Type": "Weekly site inspection", "Location": "Site",
        "Conducted By": "M. Lee", "Score": 92, "Findings": "Housekeeping around laydown area", "Status": "Open"}])

reg("Env Monitoring", "Parameter", [
    P(), Col("Parameter", "text", 20), Col("Location", "text", 18), Col("Monitoring Date", "date", 13),
    Col("Reading", "dec", 10), Col("Unit", "text", 8), Col("Limit", "dec", 10), Col("Limit Type", "list", 10, "LimitType"),
    Col("Corrective Action", "long", 30), Col("Notes", "long", 24),
    Col("Result", "calc", 12, f=lambda c, r: (
        f'=IF(OR({c("Parameter")}="",{c("Limit")}=""),"",IF({c("Limit Type")}="Min",IF({c("Reading")}<{c("Limit")},"Exceedance","OK"),'
        f'IF({c("Reading")}>{c("Limit")},"Exceedance","OK")))')),
], ex=[{"Project": EX_PROJ, "Parameter": "Turbidity", "Location": "Downstream of works", "Monitoring Date": d(-5),
        "Reading": 18, "Unit": "NTU", "Limit": 25, "Limit Type": "Max"}])

reg("Site Visits", "Location", [
    P(), Col("Visit Date", "date", 13), Col("Location", "text", 22), Col("Visited By", "text", 16),
    Col("Notes", "long", 34), Col("Discussion Notes", "long", 34),
    Col("Photos Link", "link", 30, note="Link to the OneDrive/SharePoint folder of photos for this visit."),
], ex=[{"Project": EX_PROJ, "Visit Date": d(-3), "Location": "Pump station site", "Visited By": "J. Smith",
        "Notes": "Footings complete, pumps on site next week"}])

reg("Training Records", "Person Name", [
    P(), Col("Person Name", "text", 22), Col("Company", "text", 20), Col("Training Type", "text", 22),
    Col("Completion Date", "date", 14), Col("Expiry Date", "date", 13), Col("Certificate Number", "text", 16),
    Col("Notes", "long", 24),
    Col("Expiry Flag", "calc", 14, f=lambda c, r: f'=IF({c("Person Name")}="","",{expiry(c, "Expiry Date")})'),
], ex=[{"Project": EX_PROJ, "Person Name": "T. Nguyen", "Company": "Sparky Electrical", "Training Type": "Confined space entry",
        "Completion Date": d(-300), "Expiry Date": d(65), "Certificate Number": "CS-7781"}])

reg("Toolbox Talks", "Topic", [
    P(), Col("Talk Date", "date", 13), Col("Topic", "text", 30), Col("Conducted By", "text", 18),
    Col("Attendees", "long", 28), Col("Key Points", "long", 34), Col("Notes", "long", 24),
], ex=[{"Project": EX_PROJ, "Talk Date": d(-2), "Topic": "Working near live electrical", "Conducted By": "M. Lee",
        "Attendees": "Site crew (6)", "Key Points": "Isolation and lock-out procedure"}])

reg("Lessons Learned", "Title", [
    P(), Col("Type", "list", 11, "LessonType"), Col("Title", "text", 32), Col("Category", "text", 16),
    Col("Description", "long", 34), Col("Recommendation", "long", 34),
], ex=[{"Project": EX_PROJ, "Type": "Negative", "Title": "Long-lead switchboard ordered late", "Category": "Procurement",
        "Description": "Order placed after design freeze", "Recommendation": "Order long-lead items at 60% design"}])

reg("Enquiries", "Description", [
    P(), Col("Reference", "text", 12), Col("Category", "list", 11, "EnquiryCategory"), Col("Source", "text", 14),
    Col("Contact Name", "text", 18), Col("Contact Details", "text", 22), Col("Description", "long", 34),
    Col("Date Received", "date", 13), Col("Status", "list", 11, "EnquiryStatus"), Col("Assigned To", "text", 16),
    Col("Resolution", "long", 28), Col("Date Resolved", "date", 13), Col("Notes", "long", 22),
], ex=[{"Project": EX_PROJ, "Reference": "ENQ-01", "Category": "Noise", "Source": "Phone", "Contact Name": "Resident, Main St",
        "Description": "Early morning pump testing noise", "Date Received": d(-6), "Status": "In Progress", "Assigned To": "J. Smith"}])

reg("Tenders", "Title", [
    P(), Col("Package Number", "text", 14), Col("Title", "text", 30), Col("Description", "long", 30),
    Col("Method", "list", 18, "TenderMethod"), Col("Status", "list", 11, "TenderStatus"), Col("Date Issued", "date", 13),
    Col("Closing Date", "date", 13), Col("Submissions", "int", 11), Col("Estimated Value", "money", 14),
    Col("Awarded To", "text", 20), Col("Awarded Value", "money", 14), Col("Date Awarded", "date", 13), Col("Notes", "long", 22),
], ex=[{"Project": EX_PROJ, "Package Number": "T-2201", "Title": "Electrical works package", "Method": "Selective",
        "Status": "Awarded", "Date Issued": d(-120), "Closing Date": d(-100), "Submissions": 4, "Estimated Value": 90000,
        "Awarded To": "Sparky Electrical", "Awarded Value": 85000, "Date Awarded": d(-90)}])

reg("Performance Reviews", "Reviewee Name", [
    P(), Col("Reviewee Name", "text", 24), Col("Review Period", "text", 14), Col("Review Date", "date", 13),
    Col("Quality", "list", 9, "Rating"), Col("Safety", "list", 9, "Rating"), Col("Timeliness", "list", 10, "Rating"),
    Col("Communication", "list", 13, "Rating"), Col("Comments", "long", 34),
    Col("Overall (1-5)", "calc", 11, f=lambda c, r: (
        f'=IF(OR({c("Reviewee Name")}="",COUNT({c("Quality")},{c("Safety")},{c("Timeliness")},{c("Communication")})<4),"",'
        f'AVERAGE({c("Quality")},{c("Safety")},{c("Timeliness")},{c("Communication")}))'), fmt="0.0"),
    Col("Band", "calc", 12, f=lambda c, r: (
        f'=IF({c("Overall (1-5)")}="","",IF({c("Overall (1-5)")}<2.5,"Low",IF({c("Overall (1-5)")}<3.5,"Fair","Good")))')),
], ex=[{"Project": EX_PROJ, "Reviewee Name": "ABC Civil Pty Ltd", "Review Period": "Q3", "Review Date": d(-10),
        "Quality": 4, "Safety": 4, "Timeliness": 3, "Communication": 5, "Comments": "Good workmanship; programme slipping"}])

reg("Checklist", "Title", [
    P(), Col("Phase", "list", 15, "ChecklistPhase"), Col("Title", "text", 32), Col("Description", "long", 30),
    Col("Responsible", "text", 16), Col("Due Date", "date", 13), Col("Completed Date", "date", 13),
    Col("Status", "list", 14, "ChecklistStatus"), Col("Notes", "long", 24),
    Col("Flag", "calc", 12, f=lambda c, r: (
        f'=IF(OR({c("Title")}="",{c("Status")}="Complete",{c("Status")}="Not Applicable",{c("Due Date")}=""),"",IF({c("Due Date")}<TODAY(),"Overdue",""))')),
], ex=[{"Project": EX_PROJ, "Phase": "Commissioning", "Title": "Pump performance test", "Responsible": "ABC Civil",
        "Due Date": d(45), "Status": "Not Started"}])

reg("Activity Log", "Timestamp", [
    Col("Timestamp", "date", 20), Col("User", "text", 18), Col("Sheet", "text", 20), Col("Cell", "text", 10),
    Col("Old Value", "text", 30), Col("New Value", "text", 30),
], cap=5000, blurb="Written by the macro (ThisWorkbook module). Leave empty if macros are not installed.")


# ---------------------------------------------------------------------------
# Rendering helpers
# ---------------------------------------------------------------------------
def style_header(cell, calc=False):
    cell.font = HDR_FONT
    cell.fill = HDR_CALC_FILL if calc else HDR_FILL
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = BORDER


def fmt_for(col):
    return {"date": DATE_FMT, "money": MONEY_FMT, "int": "0", "dec": "0.00"}.get(col.kind, col.fmt)


def add_rag_cf(ws, rng):
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"Green"'], fill=GREEN_FILL))
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"Amber"'], fill=AMBER_FILL))
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"Red"'], fill=RED_FILL))


def add_flag_cf(ws, rng):
    for text, fill in (("Overdue", RED_FILL), ("Expired", RED_FILL), ("Exceedance", RED_FILL), ("Critical", RED_FILL),
                       ("Due Soon", AMBER_FILL), ("Expiring Soon", AMBER_FILL), ("Due Today", AMBER_FILL),
                       ("Due This Week", AMBER_FILL), ("High", AMBER_FILL), ("Done", GREY_FILL)):
        ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=[f'"{text}"'], fill=fill))


def build_register(wb, rg, with_examples):
    ws = wb[rg.name] if rg.name in wb.sheetnames else wb.create_sheet(rg.name)
    ws.sheet_view.zoomScale = 90
    for col in rg.cols:
        cell = ws[f"{col.letter}1"]
        cell.value = col.h
        style_header(cell, col.is_calc)
        ws.column_dimensions[col.letter].width = col.w
        if col.note:
            from openpyxl.comments import Comment
            cell.comment = Comment(col.note, "APD Tracker")
    ws.row_dimensions[1].height = 34
    ws.freeze_panes = "B2" if rg.name != "Activity Log" else "A2"
    ws.auto_filter.ref = f"A1:{rg.last}{rg.cap + 1}"

    examples = rg.ex if with_examples else []
    for r in range(2, rg.cap + 2):
        c = make_c(rg.name, r)
        for col in rg.cols:
            cell = ws[f"{col.letter}{r}"]
            nf = fmt_for(col)
            if nf:
                cell.number_format = nf
            if col.is_calc:
                cell.value = col.f(c, r)
                cell.font = CALC_FONT
                cell.fill = CALC_FILL
            else:
                cell.font = INPUT_FONT
                if col.kind == "long":
                    cell.alignment = Alignment(wrap_text=True, vertical="top")
        if r - 2 < len(examples):
            for h, v in examples[r - 2].items():
                ws[f"{rg.map[h]}{r}"].value = v

    # data validation -------------------------------------------------------
    for col in rg.cols:
        rng = f"{col.letter}2:{col.letter}{rg.cap + 1}"
        if col.kind == "list" and col.lst:
            dv = DataValidation(type="list", formula1=f"=lst_{col.lst}", allow_blank=True)
            dv.error, dv.errorTitle = "Pick a value from the list.", "Invalid entry"
            ws.add_data_validation(dv); dv.add(rng)
        elif col.kind == "proj":
            dv = DataValidation(type="list", formula1="=lst_Projects", allow_blank=True)
            dv.error, dv.errorTitle = "Pick a project from the Projects sheet (add it there first).", "Unknown project"
            ws.add_data_validation(dv); dv.add(rng)
        elif col.kind == "date":
            dv = DataValidation(type="date", operator="greaterThan", formula1="36526", allow_blank=True)
            dv.error, dv.errorTitle = "Enter a valid date (e.g. 31-Dec-2026).", "Invalid date"
            ws.add_data_validation(dv); dv.add(rng)
        elif col.kind in ("money", "int", "dec"):
            dv = DataValidation(type="decimal", operator="between", formula1="-1000000000", formula2="1000000000", allow_blank=True)
            dv.error, dv.errorTitle = "Enter a number.", "Invalid number"
            ws.add_data_validation(dv); dv.add(rng)
    if rg.name == "Projects":
        dv = DataValidation(type="custom", formula1="=COUNTIF($A$2:$A$101,A2)=1", allow_blank=True)
        dv.error, dv.errorTitle = "Project names must be unique.", "Duplicate project"
        ws.add_data_validation(dv); dv.add("A2:A101")

    # conditional formatting -------------------------------------------------
    for col in rg.cols:
        rng = f"{col.letter}2:{col.letter}{rg.cap + 1}"
        if col.h in ("Schedule Health", "Budget Health", "Risk Health", "Overall Health"):
            add_rag_cf(ws, rng)
        elif col.h in ("Flag", "Overdue", "Bucket", "Expiry Flag", "Insurance Flag", "Result", "Inherent Band", "Residual Band"):
            add_flag_cf(ws, rng)
        elif col.h in ("Status",):
            ws.conditional_formatting.add(rng, FormulaRule(formula=[f'OR({col.letter}2="Done",{col.letter}2="Closed",{col.letter}2="Complete",{col.letter}2="Paid")'],
                                                           font=Font(name=FONT, color="7F7F7F")))
    return ws


def make_defined_names(wb):
    def add(name, ref):
        wb.defined_names[name] = DefinedName(name, attr_text=ref)

    lists = wb["Lists"]
    for i, (name, values) in enumerate(LISTS.items(), start=1):
        lists.cell(row=1, column=i, value=name).font = HDR_FONT
        lists.cell(row=1, column=i).fill = HDR_FILL
        lists.column_dimensions[L(i)].width = max(14, len(name) + 2)
        for j, v in enumerate(values, start=2):
            lists.cell(row=j, column=i, value=v).font = Font(name=FONT, size=10)
        add(f"lst_{name}", f"Lists!${L(i)}$2:${L(i)}${len(values) + 1}")
    add("lst_Projects", "OFFSET(Projects!$A$2,0,0,MAX(1,COUNTA(Projects!$A$2:$A$101)),1)")


# ---------------------------------------------------------------------------
# Presentation sheets
# ---------------------------------------------------------------------------
ARIAL = lambda **kw: Font(name=FONT, size=kw.pop("size", 10), **kw)
SEC_FILL = PatternFill("solid", fgColor="D9E1F2")
CARD_FILL = PatternFill("solid", fgColor="F2F2F2")
PRJ = REGS["Projects"]


def IX(rng, m):
    """INDEX that returns "" (not 0) for an empty source cell, and "" on lookup failure."""
    return f'IFERROR(IF(INDEX({rng},{m})="","",INDEX({rng},{m})),"")'


def PL(h):  # absolute Projects column range
    return R("Projects", h)


def title(ws, text, sub=None, width_cols=8):
    ws["A1"] = text
    ws["A1"].font = Font(name=FONT, size=16, bold=True, color=NAVY)
    ws.row_dimensions[1].height = 26
    if sub:
        ws["A2"] = sub
        ws["A2"].font = Font(name=FONT, size=9, italic=True, color="595959")


def section(ws, cell, text, span=1):
    ws[cell] = text
    ws[cell].font = Font(name=FONT, size=11, bold=True, color=NAVY)
    r = ws[cell].row
    c0 = ws[cell].column
    for c in range(c0, c0 + span):
        ws.cell(row=r, column=c).fill = SEC_FILL


def hdr_row(ws, row, col0, labels):
    for i, lab in enumerate(labels):
        cell = ws.cell(row=row, column=col0 + i, value=lab)
        style_header(cell)


def body(cell, fmt=None, bold=False, align=None):
    cell.font = Font(name=FONT, size=10, bold=bold)
    cell.border = BORDER
    if fmt:
        cell.number_format = fmt
    cell.alignment = Alignment(horizontal=align, vertical="top", wrap_text=fmt is None)


def build_dashboard(ws):
    title(ws, "APD Project Management Dashboard", "Every figure here is a live formula over the register sheets.")
    ws["F1"] = "As at"
    ws["F1"].font = ARIAL(bold=True)
    ws["G1"] = "=TODAY()"
    ws["G1"].number_format = DATE_FMT
    ws["G1"].font = ARIAL()
    cards = [("A", "B", "Total Projects", f'=COUNTA({PL("Project Name")})', "0"),
             ("C", "D", "In Construction", f'=COUNTIF({PL("Status")},"Construction")', "0"),
             ("E", "F", "Overdue Projects", f'=COUNTIF({PL("Overdue")},"Overdue")', "0"),
             ("G", "H", "Total Approved Budget", f'=SUM({PL("Effective Budget")})', MONEY_FMT)]
    for c1, c2, label, f, nf in cards:
        ws.merge_cells(f"{c1}4:{c2}4"); ws.merge_cells(f"{c1}5:{c2}5")
        ws[f"{c1}4"] = label
        ws[f"{c1}4"].font = ARIAL(size=9, color="595959")
        ws[f"{c1}5"] = f
        ws[f"{c1}5"].font = Font(name=FONT, size=18, bold=True, color=NAVY)
        ws[f"{c1}5"].number_format = nf
        ws[f"{c1}5"].alignment = Alignment(horizontal="left")
        for col in (c1, c2):
            for r in (4, 5):
                ws[f"{col}{r}"].fill = CARD_FILL
        ws.row_dimensions[5].height = 30
    ws.conditional_formatting.add("E5", CellIsRule(operator="greaterThan", formula=["0"], font=Font(name=FONT, size=18, bold=True, color="C00000")))

    # Portfolio table ---------------------------------------------------------
    section(ws, "A7", "Portfolio (first 30 projects - full list on the Projects sheet)", 15)
    cols = [("Project", None, None), ("Status", "Status", None), ("Project Manager", "Project Manager", None),
            ("Revised Completion", "Revised Completion", DATE_FMT), ("Effective Budget", "Effective Budget", MONEY_FMT),
            ("Forecast at Completion", "Forecast at Completion", MONEY_FMT), ("Budget Variance", "Budget Variance", MONEY_FMT),
            ("Schedule", "Schedule Health", None), ("Budget", "Budget Health", None), ("Risk", "Risk Health", None),
            ("Overall", "Overall Health", None), ("Open Tasks", "Open Tasks", "0"), ("Open Actions", "Open Actions", "0"),
            ("Open RFIs", "Open RFIs", "0"), ("Open Defects", "Open Defects", "0")]
    hdr_row(ws, 8, 1, [c[0] for c in cols])
    ws.row_dimensions[8].height = 30
    for i in range(30):
        r = 9 + i
        ws[f"A{r}"] = f'=IF(INDEX({PL("Project Name")},ROWS(A$9:A{r}))="","",INDEX({PL("Project Name")},ROWS(A$9:A{r})))'
        body(ws[f"A{r}"])
        for j, (lab, src, nf) in enumerate(cols[1:], start=2):
            cell = ws.cell(row=r, column=j)
            match = f"MATCH($A{r},{PL('Project Name')},0)"
            cell.value = f'=IF($A{r}="","",{IX(PL(src), match)})'
            body(cell, nf, align="center" if j >= 8 else None)
    add_rag_cf(ws, "H9:K38")
    ws.conditional_formatting.add("G9:G38", CellIsRule(operator="lessThan", formula=["0"], font=Font(name=FONT, color="C00000")))
    widths = [34, 18, 18, 16, 16, 16, 16, 11, 11, 11, 11, 11, 12, 11, 12]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[L(i)].width = w

    # Alerts ------------------------------------------------------------------
    section(ws, "A41", "Alerts", 3)
    hdr_row(ws, 42, 1, ["Alert", "Count"])
    alerts = [
        ("Projects past target completion", f'=COUNTIF({PL("Overdue")},"Overdue")'),
        ("Projects with Red overall health", f'=COUNTIF({PL("Overall Health")},"Red")'),
        ("Overdue tasks", f'=COUNTIF({R("Tasks", "Bucket")},"Overdue")'),
        ("Overdue actions / decisions", f'=COUNTIF({R("Actions", "Bucket")},"Overdue")'),
        ("Overdue RFIs", f'=COUNTIF({R("RFIs", "Flag")},"Overdue")'),
        ("RFIs due within 7 days", f'=COUNTIF({R("RFIs", "Flag")},"Due Soon")'),
        ("Overdue defects", f'=COUNTIF({R("Defects", "Flag")},"Overdue")'),
        ("Overdue correspondence responses", f'=COUNTIF({R("Correspondence", "Flag")},"Overdue")'),
        ("Overdue risk actions", f'=COUNTIF({R("Risks", "Overdue")},"Overdue")'),
        ("Overdue commissioning / handover items", f'=COUNTIF({R("Checklist", "Flag")},"Overdue")'),
        ("Permits expired or expiring (60 days)", f'=COUNTIF({R("Permits", "Expiry Flag")},"Expired")+COUNTIF({R("Permits", "Expiry Flag")},"Expiring Soon")'),
        ("Insurance / warranties expired or expiring", f'=COUNTIF({R("Insurance Warranties", "Expiry Flag")},"Expired")+COUNTIF({R("Insurance Warranties", "Expiry Flag")},"Expiring Soon")'),
        ("Subcontractor insurance expired or expiring", f'=COUNTIF({R("Subcontractors", "Insurance Flag")},"Expired")+COUNTIF({R("Subcontractors", "Insurance Flag")},"Expiring Soon")'),
        ("Training records expired or expiring", f'=COUNTIF({R("Training Records", "Expiry Flag")},"Expired")+COUNTIF({R("Training Records", "Expiry Flag")},"Expiring Soon")'),
        ("Environmental exceedances", f'=COUNTIF({R("Env Monitoring", "Result")},"Exceedance")'),
        ("Open High / Critical SQE issues", f'=COUNTIFS({R("SQE Issues", "Severity")},"High",{R("SQE Issues", "Status")},"<>Closed")+COUNTIFS({R("SQE Issues", "Severity")},"Critical",{R("SQE Issues", "Status")},"<>Closed")'),
        ("Payment claims awaiting certification", f'=COUNTIF({R("Payment Claims", "Status")},"Submitted")'),
    ]
    for i, (lab, f) in enumerate(alerts):
        r = 43 + i
        ws[f"A{r}"] = lab; body(ws[f"A{r}"])
        ws[f"B{r}"] = f; body(ws[f"B{r}"], "0", align="center")
    last_alert = 42 + len(alerts)
    ws.conditional_formatting.add(f"B43:B{last_alert}", CellIsRule(operator="greaterThan", formula=["0"], fill=AMBER_FILL))

    # Today & Upcoming counts ---------------------------------------------------
    section(ws, "E41", "Today & Upcoming", 4)
    hdr_row(ws, 42, 5, ["Due", "Tasks", "Actions", "Total"])
    for i, lab in enumerate(["Overdue", "Due Today", "Due This Week", "Later", "No Due Date"]):
        r = 43 + i
        ws[f"E{r}"] = lab; body(ws[f"E{r}"])
        ws[f"F{r}"] = f'=COUNTIF({R("Tasks", "Bucket")},E{r})'; body(ws[f"F{r}"], "0", align="center")
        ws[f"G{r}"] = f'=COUNTIF({R("Actions", "Bucket")},E{r})'; body(ws[f"G{r}"], "0", align="center")
        ws[f"H{r}"] = f"=F{r}+G{r}"; body(ws[f"H{r}"], "0", align="center", bold=True)
    ws.conditional_formatting.add("F43:H43", CellIsRule(operator="greaterThan", formula=["0"], fill=RED_FILL))

    # Upcoming lists ------------------------------------------------------------
    top = last_alert + 3
    section(ws, f"A{top}", "Upcoming tasks (next 15 open, soonest first)", 4)
    hdr_row(ws, top + 1, 1, ["Task", "Project", "Due", "Bucket"])
    rk = R("Tasks", "Upcoming Rank")
    for k in range(15):
        r = top + 2 + k
        m = f'MATCH(SMALL({rk},ROWS(A${top + 2}:A{r})),{rk},0)'
        for j, (h, nf) in enumerate([("Title", None), ("Project", None), ("Due Date", DATE_FMT), ("Bucket", None)], start=1):
            ws.cell(row=r, column=j, value='=' + IX(R("Tasks", h), m))
            body(ws.cell(row=r, column=j), nf)
    add_flag_cf(ws, f"D{top + 2}:D{top + 16}")

    section(ws, f"F{top}", "Upcoming actions & decisions (next 15 open)", 5)
    hdr_row(ws, top + 1, 6, ["Action", "Project", "Owner", "Due", "Bucket"])
    rk = R("Actions", "Upcoming Rank")
    for k in range(15):
        r = top + 2 + k
        m = f'MATCH(SMALL({rk},ROWS(F${top + 2}:F{r})),{rk},0)'
        for j, (h, nf) in enumerate([("Title", None), ("Project", None), ("Owner", None), ("Due Date", DATE_FMT), ("Bucket", None)], start=6):
            ws.cell(row=r, column=j, value='=' + IX(R("Actions", h), m))
            body(ws.cell(row=r, column=j), nf)
    add_flag_cf(ws, f"J{top + 2}:J{top + 16}")

    # Top risks -----------------------------------------------------------------
    rtop = top + 19
    section(ws, f"A{rtop}", "Top open risks (highest inherent score first)", 7)
    hdr_row(ws, rtop + 1, 1, ["Risk", "Project", "Score", "Band", "Owner", "Status", "Due"])
    rk = R("Risks", "Top Rank")
    for k in range(10):
        r = rtop + 2 + k
        m = f'MATCH(LARGE({rk},ROWS(A${rtop + 2}:A{r})),{rk},0)'
        for j, (h, nf) in enumerate([("Title", None), ("Project", None), ("Inherent Score", "0"), ("Inherent Band", None),
                                     ("Owner", None), ("Status", None), ("Due Date", DATE_FMT)], start=1):
            ws.cell(row=r, column=j, value='=' + IX(R("Risks", h), m))
            body(ws.cell(row=r, column=j), nf)
    add_flag_cf(ws, f"D{rtop + 2}:D{rtop + 11}")
    ws.freeze_panes = "A4"
    ws.sheet_view.zoomScale = 90


def build_summary(ws, default_project):
    title(ws, "Project Summary", "Pick a project in C3. Print this sheet (or use the PDF macro) as a one-page project report.")
    ws["B3"] = "Project:"; ws["B3"].font = ARIAL(bold=True)
    ws.merge_cells("C3:F3")
    ws["C3"] = default_project
    ws["C3"].font = Font(name=FONT, size=12, bold=True, color="0000FF")
    ws["C3"].fill = PatternFill("solid", fgColor="FFFF99")
    dv = DataValidation(type="list", formula1="=lst_Projects", allow_blank=True)
    ws.add_data_validation(dv); dv.add("C3")
    ws.column_dimensions["A"].width = 3
    for col, w in zip("BCDEFGHI", [34, 20, 20, 22, 22, 18, 18, 18]):
        ws.column_dimensions[col].width = w
    pm = f'MATCH($C$3,{PL("Project Name")},0)'

    def pv(h):
        return "=" + IX(PL(h), pm)

    section(ws, "B5", "Overview", 3)
    for i, (lab, h, nf) in enumerate([("Status", "Status", None), ("Project Manager", "Project Manager", None),
                                      ("Project Director", "Project Director", None), ("Project Sponsor", "Project Sponsor", None),
                                      ("Location", "Location", None), ("Contractor", "Contractor", None)]):
        r = 6 + i
        ws[f"B{r}"] = lab; ws[f"B{r}"].font = ARIAL(bold=True)
        ws[f"C{r}"] = pv(h); body(ws[f"C{r}"], nf)

    section(ws, "E5", "Health", 3)
    for i, (lab, h) in enumerate([("Overall", "Overall Health"), ("Schedule", "Schedule Health"), ("Budget", "Budget Health"), ("Risk", "Risk Health")]):
        r = 6 + i
        ws[f"E{r}"] = lab; ws[f"E{r}"].font = ARIAL(bold=True)
        ws[f"F{r}"] = pv(h); body(ws[f"F{r}"], align="center")
    add_rag_cf(ws, "F6:F9")
    ws.merge_cells("E10:H12")
    ws["E10"] = pv("Health Notes")
    ws["E10"].font = ARIAL(size=9, italic=True); ws["E10"].alignment = Alignment(wrap_text=True, vertical="top")

    section(ws, "B13", "Budget", 3)
    budget_rows = [("Original Budget", "Original Budget", MONEY_FMT), ("Latest Approved Budget", "Latest Approved Budget", MONEY_FMT),
                   ("Approved Variations", "Approved Variations", MONEY_FMT), ("Effective Budget", "Effective Budget", MONEY_FMT),
                   ("Expenditure to Date", "Expenditure to Date", MONEY_FMT), ("Forecast at Completion", "Forecast at Completion", MONEY_FMT),
                   ("Variance (budget - forecast)", "Budget Variance", MONEY_FMT), ("Forecast % of budget", "Forecast % of Budget", PCT_FMT)]
    for i, (lab, h, nf) in enumerate(budget_rows):
        r = 14 + i
        ws[f"B{r}"] = lab; ws[f"B{r}"].font = ARIAL(bold=True)
        ws[f"C{r}"] = pv(h); body(ws[f"C{r}"], nf)
    ws.conditional_formatting.add("C20", CellIsRule(operator="lessThan", formula=["0"], font=Font(name=FONT, color="C00000")))

    section(ws, "E13", "Schedule", 3)
    sched = [("Start Date", "Start Date", DATE_FMT), ("Target Completion", "Target Completion", DATE_FMT),
             ("Approved EOT days", "Approved EOT Days", "0"), ("Revised Completion", "Revised Completion", DATE_FMT),
             ("Actual Completion", "Actual Completion", DATE_FMT), ("Timeframe elapsed", "Schedule Elapsed %", PCT_FMT)]
    for i, (lab, h, nf) in enumerate(sched):
        r = 14 + i
        ws[f"E{r}"] = lab; ws[f"E{r}"].font = ARIAL(bold=True)
        ws[f"F{r}"] = pv(h); body(ws[f"F{r}"], nf)

    section(ws, "E21", "Open items", 3)
    cnt = [("Open tasks", "Open Tasks"), ("Open actions", "Open Actions"), ("Open RFIs", "Open RFIs"), ("Open defects", "Open Defects")]
    for i, (lab, h) in enumerate(cnt):
        r = 22 + i
        ws[f"E{r}"] = lab; ws[f"E{r}"].font = ARIAL(bold=True)
        ws[f"F{r}"] = pv(h); body(ws[f"F{r}"], "0")

    # Milestones --------------------------------------------------------------
    section(ws, "B27", "Milestones", 6)
    hdr_row(ws, 28, 2, ["Activity", "WBS", "% Complete", "Original", "Forecast", "Slip (days)"])
    sq = R("Milestones", "Summary Seq")
    for k in range(12):
        r = 29 + k
        m = f'MATCH(ROWS(B$29:B{r}),{sq},0)'
        for j, (h, nf) in enumerate([("Task / Activity", None), ("WBS Item", None), ("% Complete", "0"),
                                     ("Original Completion", DATE_FMT), ("Forecast Completion", DATE_FMT), ("Slip (days)", "0")], start=2):
            ws.cell(row=r, column=j, value='=' + IX(R("Milestones", h), m))
            body(ws.cell(row=r, column=j), nf)
    ws.conditional_formatting.add("G29:G40", CellIsRule(operator="greaterThan", formula=["0"], font=Font(name=FONT, color="C00000")))

    # Risks ------------------------------------------------------------------
    section(ws, "B43", "Top open risks", 6)
    hdr_row(ws, 44, 2, ["Risk", "Band", "Owner", "Status", "Mitigation", "Due"])
    rk = R("Risks", "Summary Rank")
    for k in range(8):
        r = 45 + k
        m = f'MATCH(LARGE({rk},ROWS(B$45:B{r})),{rk},0)'
        for j, (h, nf) in enumerate([("Title", None), ("Inherent Band", None), ("Owner", None), ("Status", None),
                                     ("Mitigation", None), ("Due Date", DATE_FMT)], start=2):
            ws.cell(row=r, column=j, value='=' + IX(R("Risks", h), m))
            body(ws.cell(row=r, column=j), nf)
    add_flag_cf(ws, "C45:C52")

    # Variations --------------------------------------------------------------
    section(ws, "B55", "Variations", 6)
    hdr_row(ws, 56, 2, ["Title", "Reference", "Cost Impact", "Time (days)", "Status", "Date Raised"])
    sq = R("Variations", "Summary Seq")
    for k in range(8):
        r = 57 + k
        m = f'MATCH(ROWS(B$57:B{r}),{sq},0)'
        for j, (h, nf) in enumerate([("Title", None), ("Reference", None), ("Cost Impact", MONEY_FMT), ("Time Impact Days", "0"),
                                     ("Status", None), ("Date Raised", DATE_FMT)], start=2):
            ws.cell(row=r, column=j, value='=' + IX(R("Variations", h), m))
            body(ws.cell(row=r, column=j), nf)

    # Actions -----------------------------------------------------------------
    section(ws, "B67", "Open actions & decisions", 6)
    hdr_row(ws, 68, 2, ["Title", "Type", "Owner", "Due", "Status", "Bucket"])
    sq = R("Actions", "Summary Seq")
    for k in range(8):
        r = 69 + k
        m = f'MATCH(ROWS(B$69:B{r}),{sq},0)'
        for j, (h, nf) in enumerate([("Title", None), ("Type", None), ("Owner", None), ("Due Date", DATE_FMT),
                                     ("Status", None), ("Bucket", None)], start=2):
            ws.cell(row=r, column=j, value='=' + IX(R("Actions", h), m))
            body(ws.cell(row=r, column=j), nf)
    add_flag_cf(ws, "G69:G76")
    ws.print_area = "A1:I77"
    ws.page_setup.orientation = "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def build_gantt(ws):
    title(ws, "Gantt - milestones for the project selected on 'Project Summary'",
          "Blue = planned span (start to forecast completion); green = % complete. The red header marks the current week.")
    ws["A3"] = "Project:"; ws["A3"].font = ARIAL(bold=True)
    ws["B3"] = "='Project Summary'!C3"; ws["B3"].font = ARIAL(bold=True, color=NAVY)
    for col, w in zip("ABCDEF", [8, 36, 13, 13, 8, 2]):
        ws.column_dimensions[col].width = w
    hdr_row(ws, 5, 1, ["WBS", "Activity", "Start", "Finish", "%"])
    sq = R("Milestones", "Summary Seq")
    for k in range(25):
        r = 6 + k
        m = f'MATCH(ROWS(A$6:A{r}),{sq},0)'
        ws[f"A{r}"] = '=' + IX(R("Milestones", "WBS Item"), m)
        ws[f"B{r}"] = '=' + IX(R("Milestones", "Task / Activity"), m)
        ws[f"C{r}"] = '=' + IX(R("Milestones", "Start Date"), m)
        ws[f"D{r}"] = (f'=IFERROR(IF(INDEX({R("Milestones", "Forecast Completion")},{m})<>"",INDEX({R("Milestones", "Forecast Completion")},{m}),'
                       f'INDEX({R("Milestones", "Original Completion")},{m})),"")')
        ws[f"E{r}"] = '=' + IX(R("Milestones", "% Complete"), m)
        for col in "ABCDE":
            body(ws[f"{col}{r}"], DATE_FMT if col in "CD" else None)
    first_week_col = 7
    for i in range(52):
        col = L(first_week_col + i)
        ws.column_dimensions[col].width = 3.2
        cell = ws[f"{col}5"]
        if i == 0:
            cell.value = '=IF(COUNT(C6:C30)=0,TODAY()-WEEKDAY(TODAY(),2)+1,MIN(C6:C30)-WEEKDAY(MIN(C6:C30),2)+1)'
        else:
            cell.value = f"={L(first_week_col + i - 1)}5+7"
        cell.number_format = "dd-mmm"
        cell.font = Font(name=FONT, size=7, bold=True, color="FFFFFF")
        cell.fill = HDR_FILL
        cell.alignment = Alignment(text_rotation=90, horizontal="center")
    ws.row_dimensions[5].height = 42
    grid = f"{L(first_week_col)}6:{L(first_week_col + 51)}30"
    g = f"{L(first_week_col)}6"
    ws.conditional_formatting.add(grid, FormulaRule(
        formula=[f'AND($C6<>"",$D6<>"",$E6<>"",{L(first_week_col)}$5<=$C6+($D6-$C6)*$E6/100,{L(first_week_col)}$5+6>=$C6)'],
        fill=PatternFill("solid", bgColor="70AD47", fgColor="70AD47"), stopIfTrue=True))
    ws.conditional_formatting.add(grid, FormulaRule(
        formula=[f'AND($C6<>"",$D6<>"",{L(first_week_col)}$5+6>=$C6,{L(first_week_col)}$5<=$D6)'],
        fill=PatternFill("solid", bgColor="8EA9DB", fgColor="8EA9DB")))
    ws.conditional_formatting.add(f"{L(first_week_col)}5:{L(first_week_col + 51)}5", FormulaRule(
        formula=[f'AND({L(first_week_col)}$5<=TODAY(),{L(first_week_col)}$5+6>=TODAY())'],
        fill=PatternFill("solid", bgColor="C00000", fgColor="C00000")))
    ws.freeze_panes = "G6"


def build_calc_and_charts(wb):
    calc = wb["Calc"]
    calc["A1"] = "Working figures for the Charts sheet - do not edit."
    calc["A1"].font = ARIAL(italic=True)
    hdr_row(calc, 2, 1, ["Status", "Projects"])
    for i, st in enumerate(LISTS["ProjectStatus"]):
        calc.cell(row=3 + i, column=1, value=st)
        calc.cell(row=3 + i, column=2, value=f'=COUNTIF({PL("Status")},A{3 + i})')
    hdr_row(calc, 2, 4, ["Month", "Spend in month", "Cumulative", "Approved budget"])
    mk = R("Cash Flow", "Month Key")
    for i in range(24):
        r = 3 + i
        calc[f"D{r}"] = (f'=IF(COUNT({mk})=0,DATE(YEAR(TODAY()),MONTH(TODAY()),1),MIN({mk}))' if i == 0 else f"=EDATE(D{r - 1},1)")
        calc[f"D{r}"].number_format = "mmm-yy"
        calc[f"E{r}"] = f'=SUMIFS({R("Cash Flow", "Counted Amount")},{mk},D{r})'
        calc[f"F{r}"] = f"=E{r}" if i == 0 else f"=F{r - 1}+E{r}"
        calc[f"G{r}"] = f'=SUM({PL("Effective Budget")})'
        for col in "EFG":
            calc[f"{col}{r}"].number_format = MONEY_FMT
    hdr_row(calc, 2, 9, ["Open risk band", "Count"])
    for i, b in enumerate(["Critical", "High", "Medium", "Low"]):
        calc.cell(row=3 + i, column=9, value=b)
        calc.cell(row=3 + i, column=10,
                  value=f'=COUNTIFS({R("Risks", "Type")},"Risk",{R("Risks", "Status")},"<>Closed",{R("Risks", "Inherent Band")},I{3 + i})')
    for col, w in zip("ABCDEFGHIJ", [20, 10, 3, 12, 16, 16, 16, 3, 16, 10]):
        calc.column_dimensions[col].width = w

    ch = wb["Charts"]
    title(ch, "Charts", "Driven by the Dashboard portfolio table and the Calc sheet.")
    b1 = BarChart(); b1.type = "col"; b1.title = "Projects by status"; b1.legend = None
    b1.add_data(Reference(calc, min_col=2, min_row=2, max_row=11), titles_from_data=True)
    b1.set_categories(Reference(calc, min_col=1, min_row=3, max_row=11))
    b1.y_axis.title = "Projects"; b1.height, b1.width = 8, 16
    b1.y_axis.delete = False; b1.x_axis.delete = False
    ch.add_chart(b1, "A4")

    dash = wb["Dashboard"]
    b2 = BarChart(); b2.type = "col"; b2.title = "Effective budget vs forecast at completion (first 8 projects)"
    b2.add_data(Reference(dash, min_col=5, max_col=6, min_row=8, max_row=16), titles_from_data=True)
    b2.set_categories(Reference(dash, min_col=1, min_row=9, max_row=16))
    b2.y_axis.number_format = "$#,##0"; b2.height, b2.width = 8, 22
    b2.y_axis.delete = False; b2.x_axis.delete = False
    ch.add_chart(b2, "K4")

    l1 = LineChart(); l1.title = "Cumulative spend (actual, then forecast) vs total approved budget"
    l1.add_data(Reference(calc, min_col=6, max_col=7, min_row=2, max_row=26), titles_from_data=True)
    l1.set_categories(Reference(calc, min_col=4, min_row=3, max_row=26))
    l1.y_axis.number_format = "$#,##0"; l1.height, l1.width = 8, 20
    l1.y_axis.delete = False; l1.x_axis.delete = False; l1.x_axis.number_format = "mmm-yy"
    ch.add_chart(l1, "A22")

    b3 = BarChart(); b3.type = "col"; b3.title = "Open risks by band"; b3.legend = None
    b3.add_data(Reference(calc, min_col=10, min_row=2, max_row=6), titles_from_data=True)
    b3.set_categories(Reference(calc, min_col=9, min_row=3, max_row=6))
    b3.height, b3.width = 8, 16
    b3.y_axis.delete = False; b3.x_axis.delete = False
    ch.add_chart(b3, "M22")


def build_search(ws):
    title(ws, "Global Search", "Type a word in B3 and press Enter. Needs the macros (see VBA/README). Without them use Ctrl+F > Options > Within: Workbook.")
    ws["A3"] = "Search for:"; ws["A3"].font = ARIAL(bold=True)
    ws["B3"].fill = PatternFill("solid", fgColor="FFFF99"); ws["B3"].font = Font(name=FONT, size=12, bold=True, color="0000FF")
    hdr_row(ws, 5, 1, ["Sheet", "Row key", "Matching text", "Go to"])
    for col, w in zip("ABCD", [24, 34, 60, 16]):
        ws.column_dimensions[col].width = w


def build_guide(wb, with_examples):
    ws = wb["Guide"]
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 30
    ws.column_dimensions["C"].width = 100
    ws["B1"] = "APD Project Tracker"
    ws["B1"].font = Font(name=FONT, size=20, bold=True, color=NAVY)
    ws["B2"] = "A single-file, OneDrive/SharePoint-hosted version of the APD Project Management Dashboard (without the AI features)."
    ws["B2"].font = ARIAL(italic=True, color="595959")
    row = 4

    def block(head, lines):
        nonlocal row
        ws[f"B{row}"] = head
        ws[f"B{row}"].font = Font(name=FONT, size=12, bold=True, color=NAVY)
        for c in ("B", "C"):
            ws[f"{c}{row}"].fill = SEC_FILL
        row += 1
        for a, b in lines:
            ws[f"B{row}"] = a; ws[f"B{row}"].font = ARIAL(bold=True)
            ws[f"B{row}"].alignment = Alignment(vertical="top", wrap_text=True)
            ws[f"C{row}"] = b; ws[f"C{row}"].font = ARIAL()
            ws[f"C{row}"].alignment = Alignment(wrap_text=True, vertical="top")
            ws.row_dimensions[row].height = max(15, 15 * (len(b) // 95 + 1))
            row += 1
        row += 1

    block("Getting started", [
        ("1. Add a project", "Go to Projects and type the project name in the first empty row (column A). Everything else on that row is optional; grey columns calculate themselves."),
        ("2. Log records", "On any register sheet, pick the Project from the dropdown in column A and fill in the row. Dropdowns, date and number checks are built in."),
        ("3. View results", "Dashboard shows the portfolio, alerts and upcoming work. Project Summary (pick a project in C3) is a printable one-page report. Gantt draws the selected project's milestones. Charts shows the trends."),
        ("4. Example rows", "Each sheet's first data row holds an example for the 'Pump Station Refurbishment (example)' project so you can see the expected format. Delete those rows' contents before real use." if with_examples
         else "This is the blank copy: every sheet is empty and ready for your data."),
    ])
    block("Colour key", [
        ("Blue text", "Input cells - type here."),
        ("Grey cells / grey headers", "Formulas - do not overwrite (Project Health, budgets, buckets, flags, ranks)."),
        ("Yellow", "Selector cells (Project Summary C3, Search B3)."),
        ("Green / Amber / Red", "Health ratings. Schedule and Budget: green below 95%, amber 95-115%, red above 115% (same rules as the web app). Risk: red if any open Critical risk, amber if any High."),
    ])
    block("How the numbers work (same rules as the web tool)", [
        ("Effective budget", "Latest Approved Budget if entered, otherwise Original Budget plus approved Variations."),
        ("Forecast at completion", "For each project and month, an Actual replaces that month's Forecast - they are never added together. Enter one line per project, month and type."),
        ("Revised completion", "Target Completion plus days from approved EOT Claims (Days Granted, or Days Claimed if none recorded)."),
        ("Overdue project", "Target Completion has passed and the status is not Project Completion, Defects Liability, Complete or Cancelled."),
        ("Risk score", "Likelihood (1-5) x Consequence (1-5): 16+ Critical, 10-15 High, 5-9 Medium, below 5 Low."),
    ])
    block("Limits (fixed row capacity - extend by copying the last formula row down)", [
        ("Projects", "100"), ("Tasks", "1,000"), ("Cash Flow", "2,000"), ("Every other register", "500 rows each"),
        ("Dashboard portfolio table", "Shows the first 30 projects; the Projects sheet always shows all."),
    ])
    block("Macros (optional - see the VBA folder)", [
        ("What they add", "Global search across every sheet, an automatic Activity Log (who changed what, when), recurring tasks (marking a recurring task Done creates the next one), one-click PDF of the Project Summary, a dated backup copy, and an Outlook email digest of overdue / expiring items."),
        ("Without macros", "Everything above that is a formula still works - dashboard, health, alerts, summary, Gantt, charts, dropdowns. Use Ctrl+F (Within: Workbook) for search."),
    ])
    block("Keeping the data in your company's OneDrive / SharePoint", [
        ("Where to save", "Save this file in a OneDrive for Business or SharePoint document library your IT department approves. The data then lives in the company tenant under its retention, backup and access policies."),
        ("Access control", "Use the library's sharing settings: Edit for the project team, View for read-only stakeholders. Excel Version History restores earlier versions."),
        ("Co-authoring", "Several people can edit the same .xlsx at once in Excel desktop or the browser. Avoid sorting a register while others are editing it."),
        ("What is not in a spreadsheet", "Per-user logins and per-project permissions, file attachments, site photo upload/map, scheduled emails and the AI assistant. Link to files in SharePoint (link columns on Documents, Meeting Minutes and Site Visits) instead."),
    ])
    ws[f"B{row}"] = "Sheet index"
    ws[f"B{row}"].font = Font(name=FONT, size=12, bold=True, color=NAVY)
    for c in ("B", "C"):
        ws[f"{c}{row}"].fill = SEC_FILL
    row += 1
    for name in wb.sheetnames:
        if name == "Guide":
            continue
        ws[f"B{row}"] = name
        ws[f"B{row}"].hyperlink = f"#'{name}'!A1"
        ws[f"B{row}"].font = Font(name=FONT, size=10, color="0563C1", underline="single")
        blurb = REGS[name].blurb if name in REGS and REGS[name].blurb else {
            "Dashboard": "Portfolio table, alerts, today & upcoming, top risks.", "Project Summary": "One-page report for the selected project.",
            "Gantt": "Weekly bar chart of the selected project's milestones.", "Charts": "Status, budget vs forecast, S-curve, risk bands.",
            "Search": "Cross-sheet search (needs the macros).", "Lists": "Dropdown values. Edit here to change a list.",
            "Calc": "Chart working figures."}.get(name, "Register - one row per record.")
        ws[f"C{row}"] = blurb; ws[f"C{row}"].font = ARIAL()
        row += 1


GROUPS = {"Guide": "7F7F7F", "Dashboard": NAVY, "Project Summary": NAVY, "Gantt": NAVY, "Charts": NAVY, "Search": NAVY,
          "Lists": "7F7F7F", "Calc": "7F7F7F", "Activity Log": "7F7F7F", "Projects": "C00000"}
ORDER = ["Guide", "Dashboard", "Project Summary", "Gantt", "Charts", "Search", "Projects", "Tasks", "Actions", "Risks", "RFIs",
         "Defects", "Variations", "EOT Claims", "Cash Flow", "Milestones", "Payment Claims", "Contracts", "Subcontractors",
         "Permits", "Insurance Warranties", "Stakeholders", "Meeting Minutes", "Correspondence", "Documents", "SQE Issues",
         "Safety Inspections", "Env Monitoring", "Site Visits", "Training Records", "Toolbox Talks", "Lessons Learned",
         "Enquiries", "Tenders", "Performance Reviews", "Checklist", "Activity Log", "Lists", "Calc"]


def build_workbook(path, with_examples, preview=False):
    wb = Workbook()
    wb.remove(wb.active)
    for name in ORDER:
        ws = wb.create_sheet(name)
        ws.sheet_properties.tabColor = GROUPS.get(name, "2E75B6")
    make_defined_names(wb)
    REGS["Tasks"].cap = max(REGS["Tasks"].cap, 1000) if not SMALL else REGS["Tasks"].cap
    for name, rg in REGS.items():
        build_register(wb, rg, with_examples)
    build_dashboard(wb["Dashboard"])
    build_summary(wb["Project Summary"], EX_PROJ if with_examples else None)
    build_gantt(wb["Gantt"])
    build_calc_and_charts(wb)
    build_search(wb["Search"])
    build_guide(wb, with_examples)
    if preview:  # small print areas so a PDF render stays short (used only for visual checks)
        for name, rg in REGS.items():
            wb[name].print_area = f"A1:{rg.last}8"
            wb[name].page_setup.orientation = "landscape"
            wb[name].page_setup.fitToWidth = 1
            wb[name].page_setup.fitToHeight = 0
            wb[name].sheet_properties.pageSetUpPr.fitToPage = True
        for name in ("Dashboard", "Gantt", "Charts", "Guide", "Search"):
            wb[name].page_setup.orientation = "landscape"
            wb[name].page_setup.fitToWidth = 1
            wb[name].page_setup.fitToHeight = 0
            wb[name].sheet_properties.pageSetUpPr.fitToPage = True
        pass
    wb.properties.title = "APD Project Tracker"
    wb.properties.creator = "APD"
    wb.active = 0
    wb.save(path)


if __name__ == "__main__":
    out = Path(sys.argv[1] if len(sys.argv) > 1 else ".")
    out.mkdir(parents=True, exist_ok=True)
    build_workbook(out / "APD_Project_Tracker.xlsx", True)
    build_workbook(out / "APD_Project_Tracker_BLANK.xlsx", False)
    print("built")
